import io

import pytest
from docx import Document

from backend.services import groq_parser
from backend.services.resume_parser import parse_resume_file
from backend.utils.file_utils import FileParsingError, FileValidationError
from tests.conftest import RESUME, make_docx, make_pdf


# ── file extraction ────────────────────────────────────────────────────

def test_pdf_text_extraction():
    text, meta = parse_resume_file(make_pdf(RESUME), 'cv.pdf')
    assert 'Reduced API latency by 40%' in text
    assert meta['file_type'] == 'pdf'


def test_docx_text_extraction_includes_tables():
    doc = Document()
    doc.add_paragraph('Jane Doe')
    doc.add_table(rows=1, cols=1).cell(0, 0).text = 'Python, Docker'
    buf = io.BytesIO()
    doc.save(buf)

    text, meta = parse_resume_file(buf.getvalue(), 'cv.docx')
    assert 'Jane Doe' in text and 'Python, Docker' in text
    assert meta['file_type'] == 'docx'


def test_empty_docx_is_rejected():
    with pytest.raises(FileParsingError):
        parse_resume_file(make_docx(''), 'cv.docx')


def test_corrupt_pdf_with_valid_header_is_rejected():
    with pytest.raises(FileParsingError):
        parse_resume_file(b'%PDF-1.4\ngarbage that is not a pdf', 'cv.pdf')


def test_renamed_file_is_rejected():
    with pytest.raises(FileValidationError, match='does not look like a valid PDF'):
        parse_resume_file(make_docx(RESUME), 'cv.pdf')


# ── Groq response handling ─────────────────────────────────────────────

@pytest.mark.parametrize('raw,expected', [
    ('{"a": 1}', {'a': 1}),
    ('```json\n{"a": 1}\n```', {'a': 1}),     # markdown fences stripped
    ('Sure! here is the JSON', None),
])
def test_try_parse_json(raw, expected):
    assert groq_parser._try_parse_json(raw) == expected


def test_validate_jd_fills_defaults_and_fixes_types():
    out = groq_parser._validate_jd_result({'job_title': 'SDE', 'required_skills': 'Python', 'keywords': None})
    assert out['required_skills'] == []        # a string is not a valid skill list
    assert out['keywords'] == [] and out['preferred_skills'] == []
    assert out['job_title'] == 'SDE'


def test_validate_resume_coerces_duration():
    out = groq_parser._validate_resume_result({'experience': [{'duration_months': '14'}, {'duration_months': 'n/a'}]})
    assert [e['duration_months'] for e in out['experience']] == [14, 0]
    assert out['skills'] == [] and out['email'] is None


def test_parse_job_description_retries_once_on_invalid_json(monkeypatch):
    responses = iter(['not json', '{"job_title": "Backend Engineer", "required_skills": ["Python"]}'])
    calls = []
    monkeypatch.setattr(groq_parser, '_get_client', lambda: object())
    monkeypatch.setattr(groq_parser, '_call_groq', lambda c, s, u: calls.append(u) or next(responses))

    out = groq_parser.parse_job_description('Backend Engineer, Python required')
    assert out['required_skills'] == ['Python']
    assert len(calls) == 2 and calls[1].startswith('Your previous response was not valid JSON')


def test_parse_resume_raises_after_two_invalid_responses(monkeypatch):
    monkeypatch.setattr(groq_parser, '_get_client', lambda: object())
    monkeypatch.setattr(groq_parser, '_call_groq', lambda c, s, u: 'still not json')
    with pytest.raises(ValueError, match='unparseable'):
        groq_parser.parse_resume(RESUME)
